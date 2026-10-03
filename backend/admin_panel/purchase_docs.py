"""Bag purchase invoice data, invoice PDF and purchases Excel export for platform admin."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Any, Dict, Iterable, List, Optional

import billing as billing_mod

_IST = timezone(timedelta(hours=5, minutes=30))


def invoice_number_for(d: dict) -> Optional[str]:
    """Stored invoice number, or the same derived number the admin list/detail already show."""
    inv = (d.get("invoice_number") or "").strip()
    if inv or d.get("status") != "PAID":
        return inv or None
    paid = d.get("paid_at") or d.get("created_at")
    short = str(d.get("id") or "").replace("-", "")[:8].upper()
    day = ""
    try:
        if paid is not None:
            day = paid.strftime("%Y%m%d") if hasattr(paid, "strftime") else str(paid)[:10].replace("-", "")
    except Exception:
        day = ""
    return f"INV-{day or 'NA'}-{short}" if short else None


def purchase_tax_view(d: dict, shop: dict) -> Dict[str, Any]:
    """Stored purchase amounts plus the CGST/SGST/IGST split used on the tax invoice."""
    gst_pct = float(d.get("gst_percent") or 0)
    gst_amt = float(d.get("gst_amount") or 0)
    gst_parts = billing_mod.split_bag_gst(
        gst_percent=gst_pct,
        gst_amount=gst_amt,
        buyer_gstin=str(shop.get("gst_number") or ""),
        buyer_state=str(shop.get("state") or ""),
    )
    return {
        "bags": int(d.get("bags") or 0),
        "price_per_bag": float(d.get("price_per_bag") or 0),
        "base_amount": float(d.get("base_amount") or 0),
        "gst_percent": gst_pct,
        "gst_amount": gst_amt,
        "total_amount": float(d.get("total_amount") or 0),
        **gst_parts,
    }


def shop_address(shop: dict) -> str:
    parts = [shop.get("address"), shop.get("village"), shop.get("taluk"), shop.get("district"), shop.get("state")]
    return ", ".join(str(p).strip() for p in parts if p and str(p).strip())


async def build_admin_invoice(db, d: dict) -> Dict[str, Any]:
    """Admin view of one purchase invoice (persists a derived invoice number once, as before)."""
    shop = await db.shops.find_one({"id": d.get("shop_id")}, {"_id": 0, "password_hash": 0}) or {}
    settings = await db.platform_billing_settings.find_one({"id": "default"}, {"_id": 0}) or {}
    hsn = (settings.get("service_hsn_code") or "998399").strip() or "998399"
    if not (d.get("invoice_number") or "").strip() and d.get("status") == "PAID":
        inv = invoice_number_for(d)
        await db.bag_purchases.update_one({"id": d["id"]}, {"$set": {"invoice_number": inv}})
        d["invoice_number"] = inv
    calc = purchase_tax_view(d, shop)
    gst_parts = {k: calc[k] for k in calc if k not in {
        "bags", "price_per_bag", "base_amount", "gst_percent", "gst_amount", "total_amount",
    }}
    return {
        **d,
        "service_hsn_code": hsn,
        "invoice_number": d.get("invoice_number") or invoice_number_for(d),
        "seller": billing_mod.bag_invoice_seller(),
        "billing_to": {
            "shop_id": shop.get("id") or d.get("shop_id"),
            "shop_name": shop.get("shop_name") or "",
            "owner_name": shop.get("owner_name") or "",
            "username": shop.get("username") or "",
            "mobile": shop.get("mobile") or "",
            "email": shop.get("email") or "",
            "address": shop_address(shop),
            "gst_number": shop.get("gst_number") or "",
        },
        "calculation": calc,
        **gst_parts,
    }


def _as_ist(value: Any) -> Optional[datetime]:
    if value is None:
        return None
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return None
    if not isinstance(value, datetime):
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(_IST)


def purchase_date_ist(d: dict) -> Optional[datetime]:
    return _as_ist(d.get("paid_at") or d.get("created_at"))


# ---------- PDF ----------
_PDF_REPLACEMENTS = {
    "\u20b9": "Rs. ",
    "\u2014": "-",
    "\u2013": "-",
    "\u00b7": "-",
    "\u2019": "'",
    "\u2018": "'",
    "\u201c": '"',
    "\u201d": '"',
    "\u00d7": "x",
}


def _pdf_text(value: Any) -> str:
    s = "" if value is None else str(value)
    for a, b in _PDF_REPLACEMENTS.items():
        s = s.replace(a, b)
    # Core PDF fonts are Latin-1 only.
    return s.encode("latin-1", "replace").decode("latin-1")


def _money(v: Any) -> str:
    return f"Rs. {float(v or 0):,.2f}"


def _pct(v: Any) -> str:
    f = float(v or 0)
    return f"{f:g}%"


def render_invoice_pdf(inv: Dict[str, Any]) -> bytes:
    from fpdf import FPDF

    seller = inv.get("seller") or {}
    bill = inv.get("billing_to") or {}
    calc = inv.get("calculation") or {}
    inv_date = purchase_date_ist(inv)

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_title(_pdf_text(f"Tax Invoice {inv.get('invoice_number') or ''}"))
    w = pdf.w - pdf.l_margin - pdf.r_margin

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(w * 0.6, 8, _pdf_text(seller.get("brand") or seller.get("name") or "LEMON MANDI"))
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(w * 0.4, 8, "TAX INVOICE", align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    right = ["BAG BALANCE"]
    left = [seller.get("legal_name")] + list(seller.get("address_lines") or [])
    if seller.get("gstin"):
        left.append(f"GSTIN: {seller['gstin']}")
    left = [x for x in left if x]
    for i in range(max(len(left), len(right))):
        pdf.cell(w * 0.6, 4.5, _pdf_text(left[i] if i < len(left) else ""))
        pdf.cell(w * 0.4, 4.5, _pdf_text(right[i] if i < len(right) else ""), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_draw_color(17, 17, 17)
    pdf.line(pdf.l_margin, pdf.get_y(), pdf.l_margin + w, pdf.get_y())
    pdf.ln(3)

    meta = [
        ("Invoice No.", inv.get("invoice_number") or "-"),
        ("Invoice Date", inv_date.strftime("%d-%m-%Y") if inv_date else "-"),
        ("Status", inv.get("status") or "-"),
        ("HSN/SAC", inv.get("service_hsn_code") or "-"),
        ("Supply", f"{calc.get('gst_supply_type') or '-'} (Place of supply code {calc.get('place_of_supply_state_code') or '-'})"),
    ]
    bill_lines = [
        f"{bill.get('shop_name') or '-'}" + (f" (@{bill['username']})" if bill.get("username") else ""),
        f"Owner: {bill['owner_name']}" if bill.get("owner_name") else "",
        bill.get("address") or "",
        bill.get("mobile") or "",
        bill.get("email") or "",
        f"GSTIN: {bill['gst_number']}" if bill.get("gst_number") else "GSTIN: Not provided",
    ]
    bill_lines = [x for x in bill_lines if x]
    top = pdf.get_y()
    pdf.set_font("Helvetica", "B", 9)
    pdf.cell(w * 0.5, 5, "BILL TO", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    for line in bill_lines:
        pdf.multi_cell(w * 0.5 - 4, 4.5, _pdf_text(line), new_x="LMARGIN", new_y="NEXT")
    left_bottom = pdf.get_y()
    pdf.set_xy(pdf.l_margin + w * 0.5, top)
    for label, value in meta:
        pdf.set_x(pdf.l_margin + w * 0.5)
        pdf.set_font("Helvetica", "B", 9)
        pdf.cell(w * 0.17, 5, _pdf_text(label))
        pdf.set_font("Helvetica", "", 9)
        pdf.multi_cell(w * 0.33, 5, _pdf_text(value), new_x="LMARGIN", new_y="NEXT")
    pdf.set_y(max(left_bottom, pdf.get_y()) + 5)

    cols = [("#", 0.06, "C"), ("Description", 0.40, "L"), ("HSN/SAC", 0.13, "C"), ("Bags", 0.11, "R"), ("Rate", 0.13, "R"), ("Taxable Value", 0.17, "R")]
    pdf.set_fill_color(17, 17, 17)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 9)
    for title, frac, align in cols:
        pdf.cell(w * frac, 7, title, border=1, align=align, fill=True)
    pdf.ln()
    pdf.set_text_color(17, 17, 17)
    pdf.set_font("Helvetica", "", 9)
    row = [
        "1",
        f"Prepaid bag balance - {int(calc.get('bags') or 0)} bags",
        inv.get("service_hsn_code") or "-",
        f"{int(calc.get('bags') or 0):,}",
        _money(calc.get("price_per_bag")),
        _money(calc.get("base_amount")),
    ]
    for (title, frac, align), value in zip(cols, row):
        pdf.cell(w * frac, 7, _pdf_text(value), border=1, align=align)
    pdf.ln(10)

    totals: List[tuple] = [("Taxable Amount", calc.get("base_amount"))]
    if (calc.get("gst_supply_type") or "").upper() == "INTER":
        totals.append((f"IGST ({_pct(calc.get('igst_percent'))})", calc.get("igst_amount")))
    else:
        totals.append((f"CGST ({_pct(calc.get('cgst_percent'))})", calc.get("cgst_amount")))
        totals.append((f"SGST ({_pct(calc.get('sgst_percent'))})", calc.get("sgst_amount")))
    totals.append((f"Total GST ({_pct(calc.get('gst_percent'))})", calc.get("gst_amount")))
    label_w, value_w = w * 0.30, w * 0.20
    for label, value in totals:
        pdf.set_x(pdf.l_margin + w - label_w - value_w)
        pdf.cell(label_w, 6, _pdf_text(label), border=1)
        pdf.cell(value_w, 6, _pdf_text(_money(value)), border=1, align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.set_x(pdf.l_margin + w - label_w - value_w)
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(label_w, 8, "GRAND TOTAL", border=1, fill=False)
    pdf.cell(value_w, 8, _pdf_text(_money(calc.get("total_amount"))), border=1, align="R", new_x="LMARGIN", new_y="NEXT")

    pdf.ln(6)
    pdf.set_font("Helvetica", "", 8)
    foot = []
    if inv.get("payment_ref"):
        foot.append(f"Payment ref: {inv['payment_ref']}")
    foot.append(f"Purchase ID: {inv.get('id') or '-'}  |  Shop ID: {inv.get('shop_id') or '-'}")
    foot.append("This is a computer-generated invoice.")
    for line in foot:
        pdf.multi_cell(w, 4.5, _pdf_text(line), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


# ---------- Excel ----------
EXPORT_HEADERS = [
    "Merchant Name",
    "Date of Purchase",
    "GST No.",
    "No. of Bags",
    "Rate",
    "Total",
    "CGST",
    "IGST",
    "Grand Total",
    "Transaction Reference No.",
]


def export_row(d: dict, shop: dict, shop_name: Optional[str]) -> List[Any]:
    calc = purchase_tax_view(d, shop)
    when = purchase_date_ist(d)
    return [
        shop_name or shop.get("shop_name") or d.get("shop_id") or "",
        when.replace(tzinfo=None) if when else None,
        (shop.get("gst_number") or "").strip(),
        calc["bags"],
        calc["price_per_bag"],
        calc["base_amount"],
        calc["cgst_amount"],
        calc["igst_amount"],
        calc["total_amount"],
        str(d.get("payment_ref") or "").strip() or None,
    ]


def render_purchases_xlsx(rows: Iterable[List[Any]], *, title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Purchases"
    ws.append(EXPORT_HEADERS)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="111111")
    thin = Side(style="thin", color="999999")
    for c in ws[1]:
        c.font = head_font
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=thin)
    ws.row_dimensions[1].height = 22

    n = 0
    for r in rows:
        ws.append(r)
        n += 1
    money_cols = ["E", "F", "G", "H", "I"]
    for i in range(2, n + 2):
        ws[f"B{i}"].number_format = "DD-MM-YYYY"
        ws[f"D{i}"].number_format = "#,##0"
        for col in money_cols:
            ws[f"{col}{i}"].number_format = "#,##0.00"

    if n:
        t = n + 2
        ws[f"A{t}"] = "TOTAL"
        ws[f"D{t}"] = f"=SUM(D2:D{n + 1})"
        ws[f"D{t}"].number_format = "#,##0"
        for col in ["F", "G", "H", "I"]:
            ws[f"{col}{t}"] = f"=SUM({col}2:{col}{n + 1})"
            ws[f"{col}{t}"].number_format = "#,##0.00"
        for c in ws[t]:
            c.font = Font(bold=True)
            c.border = Border(top=Side(style="thin", color="111111"))
        ws.auto_filter.ref = f"A1:J{n + 1}"

    widths = [32, 21, 20, 15, 10, 14, 12, 12, 17, 33]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    wb.properties.title = title

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
